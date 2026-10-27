import io
import re
from datetime import datetime
import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import plotly.graph_objects as go
import streamlit as st

# ---------------------------------------------------------
# Page Configuration
# ---------------------------------------------------------
st.set_page_config(
    page_title="ERP Sales Order Tracker",
    page_icon="📦",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ---------------------------------------------------------
# Comprehensive Synonym Dictionary
# ---------------------------------------------------------
STANDARD_COL_SYNONYMS = {
    "order_no": [
        "order no", "order_no", "so no", "so number", "sales order no", "order #",
        "so #", "voucher no", "voucher number", "doc no", "document no", "order id",
        "trans no", "transaction no", "reference no", "ref no", "order", "so"
    ],
    "invoice_no": [
        "invoice no", "invoice_no", "inv no", "inv number", "sales invoice no",
        "invoice #", "bill no", "voucher no", "voucher number", "doc no",
        "inv id", "invoice", "inv"
    ],
    "date": [
        "date", "order date", "so date", "inv date", "invoice date", "posting date",
        "doc date", "voucher date", "bill date", "entry date"
    ],
    "customer_name": [
        "customer name", "customer", "party name", "buyer name", "client name",
        "party", "account name", "buyer", "customer description", "ledger name"
    ],
    "customer_po": [
        "customer po", "po no", "po number", "cust po", "cust po no", "po ref",
        "po reference", "buyer order no", "order ref", "po", "cust po date/no"
    ],
    "item_code": [
        "item code", "item no", "product code", "item id", "sku", "material code",
        "part no", "item", "product id", "stock code"
    ],
    "item_name": [
        "item name", "description", "item description", "product name",
        "particulars", "material description", "item details"
    ],
    "qty": [
        "qty", "quantity", "order qty", "invoiced qty", "billed qty", "nos",
        "total qty", "dispatched qty", "so qty", "inv qty"
    ],
    "rate": [
        "rate", "unit rate", "unit price", "price", "basic rate", "item rate"
    ],
    "taxable_amt": [
        "taxable amount", "taxable amt", "amount", "base amount", "net amount",
        "total amt", "value", "taxable value", "basic amount", "sub total"
    ],
}

def detect_and_load_excel(file_bytes, doc_type="order"):
    """
    Scans the first 25 rows of the Excel sheet to find the header row containing
    critical identifiers like item, customer, qty, rate, voucher/order/invoice.
    """
    df_raw = pd.read_excel(file_bytes, header=None, nrows=25)
    
    header_idx = None
    target_tokens = ["item", "customer", "party", "qty", "quantity", "rate", "amount", "date", "no"]

    best_match_count = 0
    for idx, row in df_raw.iterrows():
        row_str = " ".join([str(val).lower() for val in row.values if pd.notna(val)])
        matches = sum(1 for token in target_tokens if token in row_str)
        if matches > best_match_count and matches >= 2:
            best_match_count = matches
            header_idx = idx

    if header_idx is None:
        header_idx = 0

    file_bytes.seek(0)
    df = pd.read_excel(file_bytes, skiprows=header_idx)
    
    # Drop rows that are completely empty
    df = df.dropna(how="all").reset_index(drop=True)
    return df, header_idx

def normalize_columns(df, doc_type="order"):
    """Maps arbitrary ERP column labels to clean standardized column names without duplicate collision."""
    # 1. Deduplicate raw columns first (keeps the first occurrence)
    df = df.loc[:, ~df.columns.duplicated(keep="first")].copy()
    
    col_mapping = {}
    used_targets = set()
    id_key = "order_no" if doc_type == "order" else "invoice_no"
    clean_cols = {col: re.sub(r'[^a-zA-Z0-9]', ' ', str(col)).strip().lower() for col in df.columns}

    # 2. Match Document ID first
    for orig_col, cleaned in clean_cols.items():
        if any(synonym == cleaned or synonym in cleaned for synonym in STANDARD_COL_SYNONYMS[id_key]):
            col_mapping[orig_col] = id_key
            used_targets.add(id_key)
            break

    # 3. Match all other standard columns (1-to-1 mapping only)
    for orig_col, cleaned in clean_cols.items():
        if orig_col in col_mapping:
            continue
        for std_key, synonyms in STANDARD_COL_SYNONYMS.items():
            if std_key in used_targets:
                continue
            matched = False
            for synonym in synonyms:
                if synonym == cleaned or synonym in cleaned:
                    col_mapping[orig_col] = std_key
                    used_targets.add(std_key)
                    matched = True
                    break
            if matched:
                break

    df = df.rename(columns=col_mapping)
    # Deduplicate again in case of any overlaps
    df = df.loc[:, ~df.columns.duplicated(keep="first")].copy()
    return df

def clean_reconciliation_data(df, doc_type="order"):
    """Guarantees single Series access and clean data types without crashing."""
    # Ensure no duplicate column names exist
    df = df.loc[:, ~df.columns.duplicated(keep="first")].copy()

    doc_id_col = "order_no" if doc_type == "order" else "invoice_no"
    
    # 1. Document ID safeguard
    if doc_id_col not in df.columns:
        candidates = [c for c in df.columns if any(k in str(c).lower() for k in ["no", "num", "id", "doc", "vouch", "code"])]
        if candidates:
            first_cand = df[candidates[0]]
            df[doc_id_col] = first_cand.iloc[:, 0] if isinstance(first_cand, pd.DataFrame) else first_cand
        else:
            prefix = "SO" if doc_type == "order" else "INV"
            df[doc_id_col] = [f"{prefix}-{i+1:04d}" for i in range(len(df))]
    else:
        if isinstance(df[doc_id_col], pd.DataFrame):
            df[doc_id_col] = df[doc_id_col].iloc[:, 0]

    # 2. Text fields safeguard (ensure single Series before .str operations)
    for col in ["customer_name", "customer_po", "item_code"]:
        if col not in df.columns:
            df[col] = "UNKNOWN"
        else:
            col_data = df[col]
            if isinstance(col_data, pd.DataFrame):
                col_data = col_data.iloc[:, 0]
            df[col] = col_data.fillna("UNKNOWN").astype(str).str.strip().str.upper()

    # Item name safeguard
    if "item_name" not in df.columns:
        df["item_name"] = df["item_code"]
    else:
        item_data = df["item_name"]
        if isinstance(item_data, pd.DataFrame):
            item_data = item_data.iloc[:, 0]
        df["item_name"] = item_data.fillna(df["item_code"]).astype(str).str.strip()

    # Date safeguard
    if "date" in df.columns:
        date_data = df["date"]
        if isinstance(date_data, pd.DataFrame):
            date_data = date_data.iloc[:, 0]
        df["date"] = pd.to_datetime(date_data, errors="coerce")
    else:
        df["date"] = pd.NaT

    # 3. Numeric fields safeguard
    for num_col in ["qty", "taxable_amt", "rate"]:
        if num_col in df.columns:
            val_data = df[num_col]
            if isinstance(val_data, pd.DataFrame):
                val_data = val_data.iloc[:, 0]
            df[num_col] = pd.to_numeric(val_data, errors="coerce").fillna(0.0)
        else:
            df[num_col] = 0.0

    # Backfill missing rate or taxable amount (Excluding GST)
    mask_rate_zero = (df["rate"] == 0) & (df["qty"] > 0) & (df["taxable_amt"] > 0)
    df.loc[mask_rate_zero, "rate"] = df["taxable_amt"] / df["qty"]

    mask_taxable_zero = (df["taxable_amt"] == 0) & (df["qty"] > 0) & (df["rate"] > 0)
    df.loc[mask_taxable_zero, "taxable_amt"] = df["qty"] * df["rate"]

    # Composite exact key: Customer Name + PO No + Item Code
    df["match_key"] = df["customer_name"].astype(str) + " || " + df["customer_po"].astype(str) + " || " + df["item_code"].astype(str)
    return df

# ---------------------------------------------------------
# Reconciliation Engine
# ---------------------------------------------------------
def perform_reconciliation(df_orders, df_invoices):
    # Ensure Document ID columns exist
    if "order_no" not in df_orders.columns:
        df_orders["order_no"] = [f"SO-{i+1:04d}" for i in range(len(df_orders))]
    if "invoice_no" not in df_invoices.columns:
        df_invoices["invoice_no"] = [f"INV-{i+1:04d}" for i in range(len(df_invoices))]

    # Group Orders by match_key
    order_grp = df_orders.groupby("match_key").agg(
        customer_name=("customer_name", "first"),
        customer_po=("customer_po", "first"),
        item_code=("item_code", "first"),
        item_name=("item_name", "first"),
        order_qty=("qty", "sum"),
        order_value=("taxable_amt", "sum"),
        earliest_order_date=("date", "min"),
        latest_order_date=("date", "max"),
        order_numbers=("order_no", lambda x: ", ".join(sorted(set(str(v) for v in x if pd.notna(v) and str(v) != 'nan')))),
    ).reset_index()

    order_grp["blended_order_rate"] = np.where(
        order_grp["order_qty"] > 0,
        order_grp["order_value"] / order_grp["order_qty"],
        0.0,
    )

    # Group Invoices by match_key
    inv_grp = df_invoices.groupby("match_key").agg(
        inv_qty=("qty", "sum"),
        inv_value=("taxable_amt", "sum"),
        earliest_inv_date=("date", "min"),
        latest_inv_date=("date", "max"),
        invoice_numbers=("invoice_no", lambda x: ", ".join(sorted(set(str(v) for v in x if pd.notna(v) and str(v) != 'nan')))),
    ).reset_index()

    inv_grp["avg_invoice_rate"] = np.where(
        inv_grp["inv_qty"] > 0,
        inv_grp["inv_value"] / inv_grp["inv_qty"],
        0.0,
    )

    # Merge on composite key
    reconciled = pd.merge(order_grp, inv_grp, on="match_key", how="outer")

    order_records = reconciled[reconciled["customer_name"].notna()].copy()
    unmatched_invoices = df_invoices[~df_invoices["match_key"].isin(order_grp["match_key"])].copy()

    order_records["inv_qty"] = order_records["inv_qty"].fillna(0.0)
    order_records["inv_value"] = order_records["inv_value"].fillna(0.0)
    order_records["avg_invoice_rate"] = order_records["avg_invoice_rate"].fillna(0.0)
    order_records["invoice_numbers"] = order_records["invoice_numbers"].fillna("-")

    # Math Calculations
    order_records["pending_qty"] = order_records["order_qty"] - order_records["inv_qty"]
    order_records["pending_value"] = np.maximum(0, order_records["pending_qty"]) * order_records["blended_order_rate"]
    order_records["rate_variance"] = np.where(
        order_records["inv_qty"] > 0,
        order_records["avg_invoice_rate"] - order_records["blended_order_rate"],
        0.0,
    )

    # Aging Buckets
    today = pd.Timestamp.now().normalize()
    order_records["age_days"] = (today - pd.to_datetime(order_records["earliest_order_date"])).dt.days.fillna(0).astype(int)

    def assign_aging(days):
        if days <= 30:
            return "0-30 Days"
        elif days <= 60:
            return "31-60 Days"
        else:
            return "60+ Days"

    order_records["aging_bucket"] = order_records["age_days"].apply(assign_aging)

    # Status Definitions
    def get_status(row):
        if row["order_qty"] < row["inv_qty"]:
            return "Over Delivered"
        elif row["pending_qty"] == 0:
            return "Fully Delivered"
        elif row["inv_qty"] > 0 and row["pending_qty"] > 0:
            return "Partially Delivered"
        else:
            return "Pending Execution"

    order_records["status"] = order_records.apply(get_status, axis=1)
    over_delivered = order_records[order_records["status"] == "Over Delivered"].copy()

    # Customer Summary
    customer_summary = order_records.groupby("customer_name").agg(
        total_order_value=("order_value", "sum"),
        total_invoiced_value=("inv_value", "sum"),
        total_pending_value=("pending_value", "sum"),
        pending_0_30=("pending_value", lambda s: s[order_records.loc[s.index, "aging_bucket"] == "0-30 Days"].sum()),
        pending_31_60=("pending_value", lambda s: s[order_records.loc[s.index, "aging_bucket"] == "31-60 Days"].sum()),
        pending_60_plus=("pending_value", lambda s: s[order_records.loc[s.index, "aging_bucket"] == "60+ Days"].sum()),
    ).reset_index()

    return order_records, customer_summary, unmatched_invoices, over_delivered

# ---------------------------------------------------------
# Excel Report Generator
# ---------------------------------------------------------
def create_excel_workbook(df_cust_summary, df_order_detail, df_unmatched, df_over_delivered):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=9)
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    alt_row_fill = PatternFill(start_color="F2F5F8", end_color="F2F5F8", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="D3D3D3"),
        right=Side(style="thin", color="D3D3D3"),
        top=Side(style="thin", color="D3D3D3"),
        bottom=Side(style="thin", color="D3D3D3"),
    )

    sheets_data = [
        ("Customer Summary", df_cust_summary),
        ("Order Detail", df_order_detail),
        ("Unmatched Invoices", df_unmatched),
        ("Over Delivered", df_over_delivered),
    ]

    for title, df in sheets_data:
        ws = wb.create_sheet(title=title)
        ws.views.sheetView[0].showGridLines = True

        headers = list(df.columns)
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header.replace("_", " ").title())
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for row_idx, row_data in enumerate(df.values, 2):
            for col_idx, value in enumerate(row_data, 1):
                if isinstance(value, pd.Timestamp):
                    value = value.strftime("%Y-%m-%d") if pd.notna(value) else ""
                elif isinstance(value, float) and (np.isnan(value) or np.isinf(value)):
                    value = 0.0

                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.font = data_font
                cell.border = thin_border

                if row_idx % 2 == 0:
                    cell.fill = alt_row_fill

                if isinstance(value, (int, float)):
                    if any(k in headers[col_idx - 1] for k in ["value", "amt", "rate"]):
                        cell.number_format = "#,##0.00"
                    elif any(k in headers[col_idx - 1] for k in ["qty", "days", "count"]):
                        cell.number_format = "#,##0"

        ws.freeze_panes = "A2"
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

# ---------------------------------------------------------
# Streamlit Interface
# ---------------------------------------------------------
st.title("📦 Arkaza ERP: Sales Order vs. Sales Invoice Tracker")
st.markdown("Reconcile Sales Orders and Invoices, track pending executions, calculate rate variances, and inspect fulfillment timelines.")

with st.sidebar:
    st.header("📂 Upload ERP Files")
    order_file = st.file_uploader("Upload Sales Order Export (.xlsx)", type=["xlsx"])
    invoice_file = st.file_uploader("Upload Sales Invoice Export (.xlsx)", type=["xlsx"])
    st.divider()
    st.caption("Scans up to first 25 rows for headers. Base rate / Taxable values evaluated (excl. GST).")

if order_file and invoice_file:
    with st.spinner("Processing & Reconciling ERP records..."):
        df_ord_raw, ord_head_idx = detect_and_load_excel(order_file, doc_type="order")
        df_inv_raw, inv_head_idx = detect_and_load_excel(invoice_file, doc_type="invoice")

        df_ord_norm = normalize_columns(df_ord_raw, doc_type="order")
        df_inv_norm = normalize_columns(df_inv_raw, doc_type="invoice")

        df_ord_clean = clean_reconciliation_data(df_ord_norm, doc_type="order")
        df_inv_clean = clean_reconciliation_data(df_inv_norm, doc_type="invoice")

        df_order_detail, df_cust_summary, df_unmatched, df_over_delivered = perform_reconciliation(
            df_ord_clean, df_inv_clean
        )

    # Optional Debug view for column inspection
    with st.expander("🛠️ Inspect Detected Columns & Headers"):
        c1, c2 = st.columns(2)
        c1.write(f"**Sales Order Columns (Header Row: {ord_head_idx + 1}):**")
        c1.json(list(df_ord_clean.columns))
        c2.write(f"**Sales Invoice Columns (Header Row: {inv_head_idx + 1}):**")
        c2.json(list(df_inv_clean.columns))

    # Top KPI Metrics
    kpi1, kpi2, kpi3, kpi4, kpi5 = st.columns(5)
    total_order_val = df_order_detail["order_value"].sum()
    total_inv_val = df_order_detail["inv_value"].sum()
    total_pend_val = df_order_detail["pending_value"].sum()

    kpi1.metric("Total Order Value", f"₹ {total_order_val:,.2f}")
    kpi2.metric("Invoiced Value", f"₹ {total_inv_val:,.2f}")
    kpi3.metric("Pending Value", f"₹ {total_pend_val:,.2f}")
    kpi4.metric("Over-Delivered Items", f"{len(df_over_delivered)}")
    kpi5.metric("Unmatched Invoices", f"{len(df_unmatched)}")

    st.markdown("---")

    # Tabs
    tab1, tab2, tab3, tab4 = st.tabs([
        "📈 Graphical Timeline",
        "📑 Customer Summary",
        "🔍 Order Detail & Reconciliation",
        "⚠️ Exceptions & Review"
    ])

    # Tab 1: Timeline
    with tab1:
        st.subheader("Order Booking vs. Invoiced Delivered Value Over Time")
        freq_option = st.radio("Frequency:", ["Daily", "Weekly", "Monthly"], horizontal=True, index=2)
        freq_map = {"Daily": "D", "Weekly": "W-MON", "Monthly": "MS"}
        period_fmt = {"Daily": "%Y-%m-%d", "Weekly": "Wk %W, %Y", "Monthly": "%b %Y"}

        valid_orders = df_ord_clean.dropna(subset=["date"]).copy()
        valid_invoices = df_inv_clean.dropna(subset=["date"]).copy()

        if not valid_orders.empty and not valid_invoices.empty:
            orders_ts = valid_orders.set_index("date").resample(freq_map[freq_option])["taxable_amt"].sum().reset_index()
            inv_ts = valid_invoices.set_index("date").resample(freq_map[freq_option])["taxable_amt"].sum().reset_index()

            orders_ts.columns = ["date", "Order Value"]
            inv_ts.columns = ["date", "Invoiced Value"]

            timeline_df = pd.merge(orders_ts, inv_ts, on="date", how="outer").fillna(0.0).sort_values("date")
            timeline_df["Period"] = timeline_df["date"].dt.strftime(period_fmt[freq_option])

            fig = go.Figure()
            fig.add_trace(go.Bar(x=timeline_df["Period"], y=timeline_df["Order Value"], name="Orders Booked", marker_color="#1f77b4"))
            fig.add_trace(go.Bar(x=timeline_df["Period"], y=timeline_df["Invoiced Value"], name="Invoiced Delivered", marker_color="#2ca02c"))
            fig.update_layout(barmode="group", xaxis_title="Period", yaxis_title="Taxable Amount (Excl. GST)", template="plotly_white", height=430)
            st.plotly_chart(fig, use_container_width=True)
        else:
            st.info("No transaction dates detected in files to draw timeline.")

    # Tab 2: Customer Summary
    with tab2:
        st.subheader("Customer Aging & Pending Balances")
        st.dataframe(
            df_cust_summary.style.format({
                "total_order_value": "₹ {:,.2f}",
                "total_invoiced_value": "₹ {:,.2f}",
                "total_pending_value": "₹ {:,.2f}",
                "pending_0_30": "₹ {:,.2f}",
                "pending_31_60": "₹ {:,.2f}",
                "pending_60_plus": "₹ {:,.2f}",
            }),
            use_container_width=True,
            hide_index=True
        )

    # Tab 3: Detailed Orders
    with tab3:
        st.subheader("Line-Item Reconciliation")
        c1, c2 = st.columns(2)
        cust_filter = c1.multiselect("Filter Customer:", sorted(df_order_detail["customer_name"].unique()))
        status_filter = c2.multiselect("Filter Status:", sorted(df_order_detail["status"].unique()))

        filtered = df_order_detail.copy()
        if cust_filter:
            filtered = filtered[filtered["customer_name"].isin(cust_filter)]
        if status_filter:
            filtered = filtered[filtered["status"].isin(status_filter)]

        disp_cols = [
            "customer_name", "customer_po", "item_code", "order_numbers", "invoice_numbers",
            "order_qty", "inv_qty", "pending_qty", "blended_order_rate", "avg_invoice_rate",
            "rate_variance", "pending_value", "aging_bucket", "status"
        ]

        st.dataframe(
            filtered[disp_cols].style.format({
                "order_qty": "{:,.2f}",
                "inv_qty": "{:,.2f}",
                "pending_qty": "{:,.2f}",
                "blended_order_rate": "₹ {:,.2f}",
                "avg_invoice_rate": "₹ {:,.2f}",
                "rate_variance": "₹ {:,.2f}",
                "pending_value": "₹ {:,.2f}",
            }),
            use_container_width=True,
            hide_index=True
        )

    # Tab 4: Exceptions
    with tab4:
        st.subheader("Audit Exceptions")
        t_over, t_unmatch = st.tabs(["Over-Delivered Items", "Unmatched Invoices"])
        with t_over:
            if not df_over_delivered.empty:
                st.dataframe(
                    df_over_delivered[["customer_name", "customer_po", "item_code", "order_qty", "inv_qty", "pending_qty", "blended_order_rate", "avg_invoice_rate"]],
                    use_container_width=True,
                    hide_index=True
                )
            else:
                st.success("No over-deliveries detected.")

        with t_unmatch:
            if not df_unmatched.empty:
                st.dataframe(
                    df_unmatched[["invoice_no", "date", "customer_name", "customer_po", "item_code", "qty", "rate", "taxable_amt"]],
                    use_container_width=True,
                    hide_index=True
                )
            else:
                st.success("All invoices match an existing sales order.")

    # Excel Download
    st.markdown("---")
    excel_file = create_excel_workbook(df_cust_summary, df_order_detail, df_unmatched, df_over_delivered)
    st.download_button(
        label="📥 Download Multi-Tab Reconciliation Report (.xlsx)",
        data=excel_file,
        file_name=f"Reconciliation_Report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary"
    )
else:
    st.info("👈 Please upload both Sales Order and Sales Invoice files in the left sidebar.")